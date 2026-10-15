#!/usr/bin/env python3
"""Summarise AI visibility tracking results.

Usage: python3 analyze.py <results.csv|prompt-set-*.xlsx> [--brand BYD] [--aliases "บีวายดี,byd"]
                          [--prompts prompt-set-import-BYD.csv] [--brain brand-brain-BYD.md] [--out report]
Input A: tracking export CSV from the class system, columns:
  run_date,run_label,project_brand,prompt,topic,branded,stage,platform,run_no,status,
  brand_mentioned,first_position,brands_mentioned,our_domain_cited,cited_domains,answer_text
Input B: the "Tracking (วัดด้วยมือ)" sheet of prompt-finder's Excel (needs openpyxl).
--prompts: the prompt set the run should use (prompt-finder CSV/Excel). Gives the stable IDs P01–P10.
           Without it, the earliest run in the file is the reference set.
--aliases: other names of our brand (comma or | separated) — used to check the branded label.
--brain:   path of brand-brain-<brand>.md if there is one (else the report warns that branded facts can't be verified).
Writes <out>.json (numbers for the AI to reason over) and <out>.md (tables for the user).
Only rows with status=ok count in rates; errors / no answer are reported separately.
Mention rate comes with n (answers counted) and ±95% (normal approximation, as taught in class).
Share of Voice = our brand names ÷ all brand names in the answers (each brand counted once per answer,
class definition, not position-weighted).
Prompt text is compared after Unicode NFC + whitespace clean-up (never NFKC: it splits Thai "ำ").
Warnings (wording drift, branded label vs text) go to the top of stdout, .md and .json.
"""
import csv, difflib, json, math, os, re, sys, unicodedata
from collections import Counter, defaultdict

POS_W = {1: 1.0, 2: 0.8, 3: 0.6, 4: 0.4}  # >=5 -> 0.2, absent -> 0
PLAT = {'ai_overview': 'Google AI Overviews', 'ai_mode': 'Google AI Mode', 'chatgpt': 'ChatGPT', 'gemini': 'Gemini'}
PLAT_IN = {v.lower(): k for k, v in PLAT.items()}
ORDER = list(PLAT)
FUZZY = 0.9  # difflib ratio to treat a reworded prompt as the same prompt (with a warning)
ZERO_WIDTH = dict.fromkeys(map(ord, '​‌‍⁠﻿'), None)
LABEL_Y = {'Y', 'YES', 'TRUE', '1', 'ใช่', 'BRANDED', 'B'}
LABEL_N = {'N', 'NO', 'FALSE', '0', 'ไม่', 'ไม่ใช่', 'NON-BRANDED', 'NONBRANDED', 'NON BRANDED', 'NON_BRANDED', 'NB'}


def norm(s):
    """Text key for matching prompts: NFC + no zero-width chars + single spaces + casefold (Latin only changes)."""
    s = unicodedata.normalize('NFC', str(s or '')).translate(ZERO_WIDTH)
    return ' '.join(s.split()).casefold()


def ratio(a, b):
    return difflib.SequenceMatcher(None, a, b).ratio()


def label(v):
    """Branded label → 'Y' / 'N' / None (missing or unreadable)."""
    if v is None:
        return None
    if isinstance(v, bool):
        return 'Y' if v else 'N'
    s = ' '.join(unicodedata.normalize('NFC', str(v)).split()).upper().replace('–', '-')
    return 'Y' if s in LABEL_Y else ('N' if s in LABEL_N else None)


def yn(v):
    return str(v).strip().upper() in ('Y', 'YES', 'TRUE', '1', 'ใช่')


def pos_w(p):
    try:
        p = int(p)
    except (TypeError, ValueError):
        return 0
    return POS_W.get(p, 0.2) if p > 0 else 0


def pid(v, i):
    """'P03' / 3 / '3' → 'P03'; empty → P<i>."""
    s = str(v or '').strip().upper()
    if re.fullmatch(r'P?\d{1,3}', s):
        return f'P{int(s.lstrip("P")):02d}'
    return s or f'P{i:02d}'


def _sheet_rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h or '').strip() for h in next(it)]
    return head, it


def _col(head, *names):
    for n in names:
        for i, h in enumerate(head):
            if h.lower().startswith(n.lower()):
                return i
    return None


def load_ref(path):
    """Prompt set → [{'id', 'prompt', 'branded'}] from prompt-finder CSV / Excel (or an earlier export CSV)."""
    if not path:
        return []
    ref = []
    if path.endswith('.xlsx'):
        from openpyxl import load_workbook
        head, it = _sheet_rows(load_workbook(path, read_only=True)['Prompt Set (track)'])
        ci, cp, cb = _col(head, 'ID', '#'), _col(head, 'Prompt'), _col(head, 'Branded')
        ct, cs = _col(head, 'หัวข้อ'), _col(head, 'ขั้น')
        for i, r in enumerate(it, 1):
            if r[cp]:
                ref.append({'id': pid(r[ci] if ci is not None else '', i), 'prompt': str(r[cp]),
                            'branded': label(r[cb]) if cb is not None else None,
                            'topic': r[ct] if ct is not None else '', 'stage': r[cs] if cs is not None else ''})
        return ref
    with open(path, encoding='utf-8-sig') as f:
        rows = [{(k or '').strip().lower(): v for k, v in r.items()} for r in csv.DictReader(f)]
    seen = set()
    for r in rows:
        k = norm(r.get('prompt'))
        if not k or k in seen:
            continue
        seen.add(k)
        ref.append({'id': pid(r.get('prompt_id') or r.get('id'), len(ref) + 1), 'prompt': r['prompt'],
                    'branded': label(r.get('branded')), 'topic': r.get('topic') or '', 'stage': r.get('stage') or ''})
    return ref


def load(path):
    if path.endswith('.xlsx'):
        from openpyxl import load_workbook
        head, it = _sheet_rows(load_workbook(path, read_only=True)['Tracking (วัดด้วยมือ)'])
        c = {k: _col(head, *n) for k, n in {
            'date': ('วันที่',), 'id': ('ID',), 'prompt': ('Prompt',), 'platform': ('แพลตฟอร์ม',), 'run': ('รอบที่',),
            'mention': ('ถูกเอ่ยชื่อ',), 'pos': ('ลำดับที่เอ่ย',), 'brands': ('แบรนด์ที่ถูกเอ่ย',),
            'cited_us': ('เว็บเราถูกอ้าง',), 'cited': ('เว็บที่ถูกอ้าง',)}.items()}
        g = lambda r, k: r[c[k]] if c[k] is not None and c[k] < len(r) else None
        rows = []
        for r in it:
            if g(r, 'mention') in (None, ''):
                continue  # not filled in
            rows.append({'run_date': str(g(r, 'date') or ''), 'prompt': g(r, 'prompt'), 'prompt_id': g(r, 'id') or '',
                         'platform': PLAT_IN.get(str(g(r, 'platform')).lower(), g(r, 'platform')),
                         'run_no': g(r, 'run'), 'status': 'ok', 'brand_mentioned': g(r, 'mention'),
                         'first_position': g(r, 'pos'), 'brands_mentioned': str(g(r, 'brands') or '').replace(',', '|'),
                         'our_domain_cited': g(r, 'cited_us'), 'cited_domains': str(g(r, 'cited') or '').replace(',', '|')})
        return rows
    with open(path, encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))


def assign_ids(rows, ref):
    """Give every row a stable prompt_id + the reference wording. Returns (ref, warnings)."""
    warns = []
    if not ref:
        when = lambda r: (str(r.get('run_date', '')), str(r.get('run_label', '')))
        first = min((when(r) for r in rows), default=None)
        seen = {}
        for r in rows:
            if when(r) == first and norm(r.get('prompt')) not in seen:
                seen[norm(r.get('prompt'))] = r.get('prompt')
        ref = [{'id': f'P{i:02d}', 'prompt': p, 'branded': None} for i, p in enumerate(seen.values(), 1)]
    by_id = {x['id']: x for x in ref}
    by_norm = {norm(x['prompt']): x for x in ref}
    cache, told = {}, set()
    for r in rows:
        k = norm(r.get('prompt'))
        rid = str(r.get('prompt_id') or '').strip().upper()
        if rid in by_id:
            hit = by_id[rid]
            if k != norm(hit['prompt']) and k not in told:
                told.add(k)
                warns.append(f'ถ้อยคำ prompt ไม่ตรงชุดเดิม: {hit["id"]} ในไฟล์ผลคือ "{r.get("prompt")}" '
                             f'แต่ชุดเดิมคือ "{hit["prompt"]}" — นับเป็น {hit["id"]} · ตรวจว่าระบบวัดด้วยถ้อยคำเดิมจริง')
        elif k in cache:
            hit = cache[k]
        elif k in by_norm:
            hit = by_norm[k]
        else:
            best = max(ref, key=lambda x: ratio(k, norm(x['prompt'])), default=None)
            sc = ratio(k, norm(best['prompt'])) if best else 0
            if sc >= FUZZY:
                hit = best
                warns.append(f'ถ้อยคำ prompt ไม่ตรงชุดเดิม (ตรงกัน {sc:.0%}): "{r.get("prompt")}" ≈ {best["id"]} '
                             f'"{best["prompt"]}" — นับเป็น {best["id"]} · ตรวจว่าวรรณยุกต์/ตัวสะกดหายตรงไหน '
                             f'และวัดรอบถัดไปด้วยถ้อยคำเดิม')
            else:
                hit = {'id': f'P{len(ref) + 1:02d}', 'prompt': r.get('prompt'), 'branded': None}
                ref.append(hit); by_id[hit['id']] = hit
                warns.append(f'prompt ไม่อยู่ในชุดเดิม → ให้รหัสใหม่ {hit["id"]}: "{r.get("prompt")}" '
                             f'(ถ้าแก้คำถามใน Mention Lab ถือเป็นเส้นกราฟใหม่ เทียบกับรอบก่อนไม่ได้)')
            by_norm[k] = hit
        cache[k] = hit
        r['prompt_id'], r['prompt_seen'], r['prompt'] = hit['id'], r.get('prompt'), hit['prompt']
        for f in ('topic', 'stage'):
            if not r.get(f) and hit.get(f):
                r[f] = hit[f]
    have = {r['prompt_id'] for r in rows}
    missing = [x['id'] for x in ref if x['id'] not in have]
    if not rows:
        warns.append('ไม่มีผลวัดในไฟล์นี้ (ชีต Tracking ยังไม่ได้กรอก หรือไฟล์ว่าง)')
    elif missing:
        warns.append(f'prompt ในชุดที่ไม่มีผลวัดในไฟล์นี้: {", ".join(missing)} — ถ้าถ้อยคำในระบบถูกแก้ ผลจะขึ้นเป็นรหัสใหม่แทน')
    return ref, warns


def brand_names(brand, aliases, rows):
    names = {brand, *(aliases or []), *(r.get('project_brand') for r in rows)}
    return sorted({norm(n) for n in names if n and norm(n) and norm(n) != 'เรา'}, key=len, reverse=True)


def has_brand(text, names):
    t = norm(text)
    for n in names:
        if re.fullmatch(r'[a-z0-9 .&\-]+', n):  # Latin name: whole word only ("mg" ≠ "image")
            if re.search(r'(?<![a-z0-9])' + re.escape(n) + r'(?![a-z0-9])', t):
                return True
        elif n in t:
            return True
    return False


def check_branded(rows, ref, names):
    """Decide the group of each prompt_id: 'Y' / 'N' / None (excluded until the user confirms)."""
    group, excluded, computed, warns = {}, [], [], []
    labels, have = defaultdict(set), {r['prompt_id'] for r in rows}
    for r in rows:
        lb = label(r.get('branded'))
        if lb:
            labels[r['prompt_id']].add(lb)
    for x in ref:
        if x['id'] not in have:
            continue
        lbs = labels[x['id']] or ({x['branded']} if x.get('branded') else set())
        text = has_brand(x['prompt'], names) if names else None
        if len(lbs) > 1:
            reason = 'ป้าย branded ไม่ตรงกันระหว่างรอบ (มีทั้ง Y และ N)'
        elif not lbs:
            if text is None:
                reason = 'ไม่มีป้าย branded และไม่รู้ชื่อแบรนด์ให้ตรวจ'
            else:
                group[x['id']] = 'Y' if text else 'N'
                computed.append(x['id'])
                continue
        elif text is None or (next(iter(lbs)) == 'Y') == text:
            group[x['id']] = next(iter(lbs))
            continue
        else:
            reason = ('ติดป้าย branded แต่ในข้อความไม่มีชื่อแบรนด์เรา' if 'Y' in lbs
                      else 'ติดป้าย non-branded แต่ในข้อความมีชื่อแบรนด์เรา')
        group[x['id']] = None
        excluded.append({'prompt_id': x['id'], 'prompt': x['prompt'], 'label': '/'.join(sorted(lbs)) or '-',
                         'text_has_brand': text, 'reason': reason})
    if not names:
        warns.append('ไม่ได้ตรวจป้าย branded กับข้อความ เพราะไม่รู้ชื่อแบรนด์ — รันใหม่ด้วย --brand และ --aliases')
    if computed:
        warns.append(f'ไม่มีป้าย branded ใน {len(computed)} prompt ({", ".join(computed)}) → คำนวณจากชื่อแบรนด์/ชื่อเรียกอื่นในข้อความแทน')
    if excluded:
        warns.append(f'ป้าย branded ไม่ตรงกับข้อความ {len(excluded)} prompt → ยังไม่นับในสถิติทั้ง 2 กลุ่ม '
                     f'จนกว่าจะยืนยัน (แก้ป้ายแล้วรันใหม่): ' +
                     ' · '.join(f'{e["prompt_id"]} "{e["prompt"]}" ({e["reason"]})' for e in excluded))
    return group, excluded, warns


def split(v):
    return [x.strip() for x in str(v or '').split('|') if x.strip()]


def rate(rows):
    ok = [r for r in rows if r.get('status', 'ok') == 'ok']
    return (round(100 * sum(yn(r['brand_mentioned']) for r in ok) / len(ok)) if ok else None), len(ok)


def margin(pct, n):
    """±95% in percentage points for a rate (normal approximation, e.g. 40% n=10 → ±30, n=120 → ±9)."""
    if pct is None or not n:
        return None
    p = pct / 100
    return round(196 * math.sqrt(p * (1 - p) / n))


def sov_band(v, n_brands=None):
    """Class bands (M3) for a topic with about 8–15 competing brands; fewer brands → compare with the average."""
    if v is None:
        return '-'
    if n_brands and n_brands < 8:
        avg = round(100 / n_brands, 1)
        return f'{"สูงกว่า" if v > avg else "ต่ำกว่า"}ค่าเฉลี่ย {avg}% (แบรนด์น้อย ใช้ 4 ช่วงตรง ๆ ไม่ได้)'
    return ('แทบไม่มีเสียง' if v < 5 else 'อยู่ในวง' if v < 10 else 'กลุ่มนำ' if v < 20 else 'เจ้าที่')


def sov_of(sub, key=None):
    """Share of Voice per group: {group: {brand: %}} — each brand once per answer, ÷ all brand names."""
    cnt = defaultdict(Counter)
    for r in sub:
        if r.get('status', 'ok') != 'ok':
            continue
        names = list(dict.fromkeys(split(r.get('brands_mentioned'))))
        cnt[key(r) if key else 'all'].update(names)
    return {g: {b: round(100 * v / sum(c.values()), 1) for b, v in c.most_common()} for g, c in cnt.items() if sum(c.values())}


def main(path, brand=None, out='report', aliases=None, prompts=None, brain=None):
    rows = load(path)
    brand = brand or (rows[0].get('project_brand') if rows else '') or 'เรา'
    aliases = [a.strip() for a in re.split(r'[,|]', aliases) if a.strip()] if isinstance(aliases, str) else (aliases or [])
    ref = load_ref(prompts or (path if path.endswith('.xlsx') else None))
    ref, warns = assign_ids(rows, ref)
    group, excluded, bwarns = check_branded(rows, ref, brand_names(brand, aliases, rows))
    warns += bwarns
    errors = Counter(r.get('status') for r in rows if r.get('status', 'ok') != 'ok')
    res = {'brand': brand, 'aliases': aliases, 'rows': len(rows), 'errors': dict(errors),
           'run_dates': sorted({str(r.get('run_date', '')) for r in rows}), 'warnings': warns,
           'excluded_prompts': excluded, 'prompt_ids': [{'id': x['id'], 'prompt': x['prompt']} for x in ref],
           'brand_brain': brain if brain and os.path.exists(brain) else None}

    for lbl, flag in (('non_branded', 'N'), ('branded', 'Y')):
        sub = [r for r in rows if group.get(r['prompt_id']) == flag]
        m, n = rate(sub)
        block = {'mention_rate': m, 'answers': n, 'margin': margin(m, n), 'by_platform': {}, 'by_topic': {}, 'prompts': []}
        for p in ORDER:
            pr, pn = rate([r for r in sub if r.get('platform') == p])
            if pn:
                block['by_platform'][PLAT[p]] = {'mention_rate': pr, 'answers': pn, 'margin': margin(pr, pn)}
        topic_sov = sov_of(sub, lambda r: (r.get('topic', '') or '') or '-')
        for t in sorted({r.get('topic', '') or '' for r in sub}):
            tr, tn = rate([r for r in sub if (r.get('topic', '') or '') == t])
            ts = topic_sov.get(t or '-', {})
            block['by_topic'][t or '-'] = {'mention_rate': tr, 'answers': tn, 'margin': margin(tr, tn),
                                           'sov': ts.get(brand, 0.0) if ts else None, 'brands_named': len(ts)}
        # Share of Voice (class definition: our names ÷ all brand names, once per answer) + competitor counts
        comp = Counter(b for r in sub if r.get('status', 'ok') == 'ok'
                       for b in dict.fromkeys(split(r.get('brands_mentioned'))) if b.lower() != brand.lower())
        allv = sov_of(sub).get('all', {})
        block['sov_all'] = allv.get(brand, 0.0) if allv else None
        block['brands_named'] = len(allv)
        block['sov_band'] = sov_band(block['sov_all'], len(allv))
        block['sov'] = {PLAT.get(p, p): dict(list(d.items())[:6]) for p, d in sov_of(sub, lambda r: r['platform']).items()}
        block['top_competitors'] = comp.most_common(6)
        # per prompt × platform status (matched by prompt_id, not by exact text)
        for pi in dict.fromkeys(r['prompt_id'] for r in sub):
            prs = [r for r in sub if r['prompt_id'] == pi]
            entry = {'prompt_id': pi, 'prompt': prs[0]['prompt'], 'topic': prs[0].get('topic', ''),
                     'stage': prs[0].get('stage', ''), 'platforms': {}}
            for p in ORDER:
                ps = [r for r in prs if r.get('platform') == p and r.get('status', 'ok') == 'ok']
                if not ps:
                    continue
                ours = sum(yn(r['brand_mentioned']) for r in ps)
                others = Counter(b for r in ps for b in split(r.get('brands_mentioned')) if b.lower() != brand.lower())
                best = others.most_common(1)[0] if others else (None, 0)
                status = 'Absent' if ours == 0 else ('Lead' if ours >= best[1] else 'Chasing')
                posl = [int(r['first_position']) for r in ps if yn(r['brand_mentioned']) and str(r.get('first_position', '')).isdigit()]
                entry['platforms'][PLAT[p]] = {
                    'status': status, 'mentioned': f'{ours}/{len(ps)}', 'avg_position': round(sum(posl) / len(posl), 1) if posl else None,
                    'shown_instead': [b for b, _ in others.most_common(3)],
                    'our_site_cited': sum(yn(r.get('our_domain_cited')) for r in ps),
                    'cited_domains': [d for d, _ in Counter(d for r in ps for d in split(r.get('cited_domains'))).most_common(5)]}
            st = [v['status'] for v in entry['platforms'].values()]
            entry['overall'] = 'Absent' if st and all(s == 'Absent' for s in st) else ('Lead' if st and all(s == 'Lead' for s in st) else 'Mixed')
            block['prompts'].append(entry)
        res[lbl] = block

    notes = []
    if res['branded']['answers'] and not res['brand_brain']:
        notes.append('ยังไม่มี Brand Brain (ปกติในคลาส — ทำใน Workshop 4) · ข้อมูลในคำตอบ branded (ราคา ประกัน สเปก เงื่อนไข) '
                     'จึงติด "ต้องตรวจ" ไว้ก่อน ห้ามสรุปว่าถูกหรือผิด · ทำ Brand Brain แล้วรันใหม่ด้วย --brain brand-brain-<แบรนด์>.md')
    cited = Counter(d for r in rows if r.get('status', 'ok') == 'ok' for d in split(r.get('cited_domains')))
    res['top_cited_domains'] = cited.most_common(15)
    with open(out + '.json', 'w', encoding='utf-8') as f:
        res['notes'] = notes
        json.dump(res, f, ensure_ascii=False, indent=1)

    L = [f'# AI Visibility — {brand}', f'ข้อมูล: {len(rows)} คำตอบ · รอบ: {", ".join(res["run_dates"])}' +
         (f' · ไม่นับ (error/ไม่มีคำตอบ): {dict(errors)}' if errors else ''), '']
    if warns:
        L += ['## ⚠️ ต้องดูก่อนอ่านผล'] + [f'- {w}' for w in warns] + ['']
    if excluded:
        L += ['### prompt ที่ยังไม่นับ (ป้าย branded ไม่ตรงกับข้อความ)', '| ID | Prompt | ป้าย | มีชื่อแบรนด์ในข้อความ | เหตุผล |', '|---|---|---|---|---|']
        L += [f'| {e["prompt_id"]} | {e["prompt"]} | {e["label"]} | {"ใช่" if e["text_has_brand"] else ("ไม่" if e["text_has_brand"] is False else "-")} | {e["reason"]} |'
              for e in excluded] + ['']
    for lbl, name in (('non_branded', 'Non-branded — AI เลือกเราไหม'), ('branded', 'Branded — AI พูดถึงเราอย่างไร')):
        b = res[lbl]
        if not b['answers']:
            continue
        L += [f'## {name}', f'Mention rate รวม: **{b["mention_rate"]}%** ±{b["margin"]} (n = {b["answers"]} คำตอบ)']
        if b['sov_all'] is not None:
            L.append(f'Share of Voice รวม: **{b["sov_all"]}%** ({b["sov_band"]} · AI เอ่ยชื่อ {b["brands_named"]} แบรนด์ · '
                     f'ค่าเฉลี่ย 100 ÷ {b["brands_named"]} = {round(100 / b["brands_named"], 1)}%)')
        L += ['', '| แพลตฟอร์ม | Mention rate | n | SoV ของเรา | ผู้นำ SoV |', '|---|---|---|---|---|']
        for p, v in b['by_platform'].items():
            s = b['sov'].get(p, {})
            lead = next(iter(s.items()), ('-', 0))
            L.append(f'| {p} | {v["mention_rate"]}% ±{v["margin"]} | {v["answers"]} | {s.get(brand, 0)}% | {lead[0]} {lead[1]}% |')
        L += ['', '| หัวข้อ | Mention rate | n | SoV ของเรา |', '|---|---|---|---|'] + [
            f'| {t} | {v["mention_rate"]}% ±{v["margin"]} | {v["answers"]} | ' + (f'{v["sov"]}% ({sov_band(v["sov"], v["brands_named"])})' if v['sov'] is not None else '-') + ' |'
            for t, v in b['by_topic'].items()]
        L += ['', 'SoV = ชื่อเรา ÷ ชื่อทุกแบรนด์ในคำตอบ (นับแบรนด์ละ 1 ครั้งต่อคำตอบ) · เกณฑ์ในคลาส (หัวข้อที่มีคู่แข่งราว 8–15 แบรนด์): '
              '<5% แทบไม่มีเสียง · 5–10% อยู่ในวง · 10–20% กลุ่มนำ · 20%+ เจ้าที่ · แบรนด์น้อย (3–5 เจ้า) เกณฑ์ขยับขึ้น · '
              '± = ช่วงคลาดเคลื่อน 95% ต่างกันน้อยกว่านี้ = ยังสูสี']
        L += ['', '| ID | Prompt | ' + ' | '.join(PLAT[p] for p in ORDER) + ' | ใครถูกแสดงแทน |', '|---|---|' + '---|' * (len(ORDER) + 1)]
        for e in b['prompts']:
            cells = [f'{e["platforms"][PLAT[p]]["status"]} {e["platforms"][PLAT[p]]["mentioned"]}' if PLAT[p] in e['platforms'] else '-' for p in ORDER]
            inst = sorted({x for v in e['platforms'].values() for x in v['shown_instead']})
            L.append(f'| {e["prompt_id"]} | {e["prompt"]} | ' + ' | '.join(cells) + f' | {", ".join(inst[:4])} |')
        L.append('')
    L += ['## เว็บที่ AI อ้างบ่อย', '| โดเมน | ครั้ง |', '|---|---|'] + [f'| {d} | {c} |' for d, c in res['top_cited_domains']]
    if notes:
        L += ['', '### หมายเหตุ'] + [f'- {n}' for n in notes]
    L += ['', '**ขั้นต่อไปในคลาส:** Workshop 3 → /site-readiness-audit (ตรวจว่า AI เข้าเว็บและอ่านได้ไหม) · Workshop 4 → /brand-brain → /ai-content-writer → /ai-citable-page-check']
    with open(out + '.md', 'w', encoding='utf-8') as f:
        f.write('\n'.join(L))
    for w in warns:
        print('WARN:', w)
    for n in notes:
        print('NOTE:', n)
    print(f'saved {out}.json + {out}.md')
    return res


if __name__ == '__main__':
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    kw = {}
    for flag in ('brand', 'out', 'aliases', 'prompts', 'brain'):
        if f'--{flag}' in a:
            i = a.index(f'--{flag}'); kw[flag] = a[i + 1]; del a[i:i + 2]
    main(a[0], **kw)
