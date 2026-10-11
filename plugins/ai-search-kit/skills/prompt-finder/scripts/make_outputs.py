#!/usr/bin/env python3
"""Turn a prompt-finder JSON file into an Excel workbook + an import CSV for the tracking system.

Usage: python3 make_outputs.py prompts.json <brand-slug> [out_dir]
Writes: prompt-set-<brand>.xlsx and prompt-set-import-<brand>.csv
Needs openpyxl for the .xlsx (pip install openpyxl). The CSV is written even without it.
The selected prompts get stable IDs P01–P10 (CSV column prompt_id, last — Mention Lab ignores extra columns).
Rules the script enforces (and warns about): source "ai" → demand = 1 · branded → mention = 3 ·
a "score" in the JSON that differs from the computed one is replaced by the computed score ·
equal scores are ordered by source: customer > gsc > keyword > user > ai.

JSON shape:
{
  "brand": "BYD", "aliases": ["BYD", "บีวายดี"], "website": "https://www.byd.com/th",
  "product": "...", "audience": "...", "competitors": ["MG", "Tesla"], "market": "TH", "language": "th",
  "topics": [{"name": "ชาร์จสะดวก", "evidence": "ลูกค้าถามบ่อยเรื่องชาร์จที่คอนโด"}],
  "prompts": [
    {"prompt": "...", "topic": "ชาร์จสะดวก", "stage": "Discovery|Evaluation|Comparison|Decision",
     "branded": false, "source": "customer|user|gsc|keyword|ai",
     "scores": {"business": 3, "demand": 2, "mention": 3, "win": 2, "risk": 2},
     "selected": true, "why": "...", "id": "P01 (optional — keep the same id when re-running)"}
  ]
}
"""
import csv, json, os, re, sys

W = {'business': .30, 'demand': .20, 'mention': .20, 'win': .15, 'risk': .15}
STAGES = ('Discovery', 'Evaluation', 'Comparison', 'Decision')
SRC = {'customer': 'ลูกค้าถามจริง', 'user': 'ผู้ใช้ให้มา', 'gsc': 'Google Search Console', 'keyword': 'keyword / AI volume', 'ai': 'AI ช่วยคิด'}
PLATFORMS = ('Google AI Overviews', 'Google AI Mode', 'ChatGPT', 'Gemini')
SRC_RANK = {'customer': 0, 'gsc': 1, 'keyword': 2, 'user': 3, 'ai': 4}  # tie-break: real evidence first


def rank(p):
    return SRC_RANK.get(p.get('source', 'ai'), 3)


def used_scores(p):
    """Scores after the fixed rules: branded → mention 3 · source ai → demand 1 (no evidence)."""
    s = dict(p.get('scores', {}))
    if p.get('branded'):
        s['mention'] = 3
    if p.get('source', 'ai') == 'ai':
        s['demand'] = 1
    return s


def score(p):
    s = used_scores(p)
    return round(sum(W[k] * s.get(k, 1) for k in W), 1)


def score_warnings(ps):
    """Call before p['score'] is overwritten: list prompts whose given score / ai-demand differ from the rules."""
    warn = []
    ai_fixed = [p for p in ps if p.get('source', 'ai') == 'ai' and p.get('scores', {}).get('demand', 1) != 1]
    if ai_fixed:
        warn.append(f'source = ai ให้ demand = 1 เสมอ (ไม่มีหลักฐานความต้องการ) — ปรับแล้ว {len(ai_fixed)} ข้อ: ' +
                    ' · '.join(f'"{p["prompt"][:40]}…" (demand {p["scores"]["demand"]} → 1)' for p in ai_fixed))
    bad = []
    for p in ps:
        try:
            given = round(float(p['score']), 1) if p.get('score') not in (None, '') else None
        except (TypeError, ValueError):
            given = p.get('score')
        if given is not None and given != score(p):
            bad.append(f'"{p["prompt"][:40]}…" ใส่มา {given} → ใช้ {score(p)}')
    if bad:
        warn.append(f'คะแนนที่ใส่มาไม่ตรงกับที่คำนวณจาก 5 ด้าน {len(bad)} ข้อ (ใช้ค่าที่คำนวณ — ถ้าตั้งใจต่าง ให้แก้คะแนนรายด้าน): ' + ' · '.join(bad))
    return warn


def ai_over_customer(ps):
    """Flag: an 'ai' prompt was selected while a 'customer' prompt with the same topic + stage + group was left out."""
    out = []
    for a in (p for p in ps if p.get('selected') and p.get('source', 'ai') == 'ai'):
        for c in ps:
            if (not c.get('selected') and c.get('source') == 'customer' and c.get('topic') == a.get('topic')
                    and c.get('stage') == a.get('stage') and bool(c.get('branded')) == bool(a.get('branded'))):
                out.append((a, c))
    return out


def check(d, chosen, ps=()):
    warn = []
    for a, c in ai_over_customer(ps):
        warn.append(f'เลือกข้อ "AI ช่วยคิด" {a.get("id", "")} "{a["prompt"][:40]}…" ({a["score"]}) แต่ข้อจากลูกค้าจริง '
                    f'"{c["prompt"][:40]}…" ({c["score"]}) หัวข้อ/ขั้นเดียวกันไม่ถูกเลือก — ตรวจว่าตั้งใจไหม (ปกติให้ข้อลูกค้ามาก่อน)')
    nb = [p for p in chosen if not p.get('branded')]
    b = [p for p in chosen if p.get('branded')]
    if len(chosen) != 10:
        warn.append(f'เลือก {len(chosen)} ข้อ (ควรเป็น 10)')
    if len(nb) != 7 or len(b) != 3:
        warn.append(f'non-branded {len(nb)} / branded {len(b)} (ควรเป็น 7 / 3)')
    if sum(p.get('stage') in ('Comparison', 'Decision') for p in nb) < 5:
        warn.append('non-branded ที่เป็น Comparison/Decision น้อยกว่า 5 ข้อ')
    if sum(p.get('stage') == 'Discovery' for p in chosen) > 1:
        warn.append('Discovery เกิน 1 ข้อ')
    names = [a.lower() for a in d.get('aliases', [])] + [d.get('brand', '').lower()]
    for p in nb:
        if any(n and n in p['prompt'].lower() for n in names):
            warn.append(f'non-branded มีชื่อแบรนด์: {p["prompt"][:40]}…')
    for p in chosen:
        if p.get('stage') not in STAGES:
            warn.append(f'stage ไม่ถูกต้อง: {p.get("stage")}')
    return warn


def write_csv(chosen, path):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['prompt', 'topic', 'branded', 'stage', 'source', 'score', 'why', 'prompt_id'])
        for p in chosen:
            w.writerow([p['prompt'], p.get('topic', ''), 'Y' if p.get('branded') else 'N', p.get('stage', ''),
                        p.get('source', 'ai'), p['score'], p.get('why', ''), p['id']])


def write_xlsx(d, ps, chosen, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    head = PatternFill('solid', fgColor='152641')
    pick = PatternFill('solid', fgColor='F7F5EC')

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = head
            c.alignment = Alignment(vertical='center', wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wd
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.freeze_panes = 'A2'

    wb = Workbook()
    ws = wb.active; ws.title = 'Prompt Set (track)'
    sheet(ws, ['ID', 'Prompt', 'หัวข้อ (เกณฑ์ตัดสินใจ)', 'Branded?', 'ขั้น', 'ที่มา', 'คะแนน (1–3)', 'ทำไมเลือก'],
          [[p['id'], p['prompt'], p.get('topic'), 'Branded' if p.get('branded') else 'Non-branded', p.get('stage'),
            SRC.get(p.get('source'), p.get('source')), p['score'], p.get('why', '')] for p in chosen],
          [4, 60, 20, 13, 12, 18, 11, 40])

    ws = wb.create_sheet('ผู้สมัครทั้งหมด')
    rows = sorted(ps, key=lambda p: (p.get('topic', ''), -p['score'], rank(p)))
    keys = list(W)
    flags = {}
    for a, c in ai_over_customer(ps):
        flags.setdefault(id(a), []).append('ข้อ AI ถูกเลือกแทนข้อลูกค้าในหัวข้อ/ขั้นเดียวกัน')
        flags.setdefault(id(c), []).append(f'ข้อลูกค้าที่ไม่ถูกเลือก (เลือก {a.get("id", "ข้อ AI")} แทน)')
    sheet(ws, ['Prompt', 'หัวข้อ', 'Branded?', 'ขั้น', 'ที่มา', 'ธุรกิจ 30%', 'ความต้องการ 20%', 'เอ่ยชื่อ 20%',
               'โอกาสชนะ 15%', 'ความเสี่ยง 15%', 'คะแนน', 'เลือก', 'ธง'],
          [[p['prompt'], p.get('topic'), 'Branded' if p.get('branded') else 'Non-branded', p.get('stage'),
            SRC.get(p.get('source'), p.get('source'))] + [used_scores(p).get(k) for k in keys] +
           [p['score'], p.get('id', '✓') if p.get('selected') else '', ' · '.join(flags.get(id(p), []))] for p in rows],
          [60, 18, 13, 12, 18, 9, 11, 9, 10, 10, 8, 7, 30])
    for row in ws.iter_rows(min_row=2):
        if row[11].value:
            for c in row:
                c.fill = pick

    ws = wb.create_sheet('Tracking (วัดด้วยมือ)')
    sheet(ws, ['วันที่', 'ID', 'Prompt', 'แพลตฟอร์ม', 'รอบที่', 'ถูกเอ่ยชื่อไหม (Y/N)', 'ลำดับที่เอ่ยครั้งแรก',
               'แบรนด์ที่ถูกเอ่ยทั้งหมด', 'เว็บเราถูกอ้างไหม (Y/N)', 'เว็บที่ถูกอ้าง', 'ข้อมูลถูกต้องไหม', 'หมายเหตุ'],
          [['', p['id'], p['prompt'], pl, r, '', '', '', '', '', '', ''] for p in chosen for pl in PLATFORMS for r in (1, 2, 3)],
          [11, 6, 50, 18, 7, 13, 12, 28, 14, 28, 14, 24])

    ws = wb.create_sheet('วิธีใช้')
    info = [
        ['แบรนด์', f"{d.get('brand', '')} ({', '.join(d.get('aliases', []))})"], ['เว็บไซต์', d.get('website', '')],
        ['สินค้า / บริการ', d.get('product', '')], ['กลุ่มลูกค้า', d.get('audience', '')],
        ['คู่แข่ง', ', '.join(d.get('competitors', []))], ['ตลาด / ภาษา', f"{d.get('market', '')} / {d.get('language', '')}"],
        ['หัวข้อที่โฟกัส', ' · '.join(t['name'] for t in d.get('topics', []))],
        ['', ''],
        ['ชีต Prompt Set', '10 ข้อที่ควร track: non-branded 7 (AI เลือกเราไหม) + branded 3 (AI พูดถึงเราถูกไหม)'],
        ['นำเข้าระบบวัดผล', 'ใช้ไฟล์ prompt-set-import-*.csv ที่ได้มาพร้อมกัน'],
        ['วัดด้วยมือ', 'ชีต Tracking: ถามแต่ละข้อใน Google AI Overviews → Google AI Mode → ChatGPT → Gemini รอบละ 3 ครั้ง (ไม่ login) แล้วกรอก · วัดซ้ำรายสัปดาห์ด้วยถ้อยคำเดิม'],
        ['ID / ถ้อยคำ', 'P01–P10 คือรหัสประจำ prompt ใช้จับคู่ผลข้ามรอบ · copy ถ้อยคำจากไฟล์นี้ไปวาง ห้ามพิมพ์ใหม่ (วรรณยุกต์/ตัวสะกดเพี้ยนตัวเดียว = prompt ใหม่)'],
        ['Mention rate', 'จำนวนคำตอบที่มีชื่อเรา ÷ คำตอบทั้งหมด (คิดแยก non-branded กับ branded) · ไม่มี "อันดับใน AI" ให้ดูแนวโน้ม'],
        ['คะแนน', '0.30×ธุรกิจ + 0.20×ความต้องการ + 0.20×เอ่ยชื่อ + 0.15×โอกาสชนะ + 0.15×ความเสี่ยงถ้า AI ตอบผิด (แต่ละด้าน 1–3) · ที่มา "AI ช่วยคิด" ความต้องการ = 1 เสมอ · คะแนนเท่ากันเรียงตามที่มา: ลูกค้า > GSC > keyword > ผู้ใช้ให้มา > AI'],
        ['ข้อควรรู้', 'prompt ไม่มี search volume จริง · ที่มา "AI ช่วยคิด" ควรเทียบกับคำถามลูกค้าจริงก่อนใช้ · ไม่ใส่ปีใน prompt'],
        ['อ้างอิงแนวคิด', 'Prompt Map · แนวคิดจาก AI SEO For Dummies (Nathan Gotch, Wiley 2026) บทที่ 5, 9 — สรุปด้วยคำของเรา'],
    ]
    for r in info:
        ws.append(r)
    ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 100
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path)


def main(src, slug, out='.'):
    d = json.load(open(src, encoding='utf-8'))
    ps = d['prompts']
    warns = score_warnings(ps)
    for p in ps:
        p['score'] = score(p)
    chosen = [p for p in ps if p.get('selected')]
    chosen.sort(key=lambda p: (bool(p.get('branded', False)), -p['score'], rank(p)))
    given = [str(p.get('id') or p.get('prompt_id') or '').strip().upper() for p in chosen]
    keep = all(re.fullmatch(r'P\d{2,3}', g) for g in given) and len(set(given)) == len(given)
    for i, p in enumerate(chosen, 1):
        p['id'] = given[i - 1] if keep else f'P{i:02d}'
    os.makedirs(out, exist_ok=True)
    csv_path = os.path.join(out, f'prompt-set-import-{slug}.csv')
    write_csv(chosen, csv_path)
    print(f'saved {csv_path}')
    try:
        xlsx_path = os.path.join(out, f'prompt-set-{slug}.xlsx')
        write_xlsx(d, ps, chosen, xlsx_path)
        print(f'saved {xlsx_path}: {len(chosen)} selected / {len(ps)} candidates')
    except ImportError:
        print('openpyxl not installed → only CSV written (pip install openpyxl for the Excel file)')
    for w in warns + check(d, chosen, ps):
        print('WARN:', w)


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(*sys.argv[1:4])
