"""Tests for the AI Search Kit scripts (analyze.py, make_outputs.py).

Run: python3 -m pytest tests/   or   python3 tests/run.py
"""
import contextlib, csv, importlib.util, io, json, os, shutil, tempfile, unicodedata

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SK = os.path.join(ROOT, 'plugins', 'ai-search-kit', 'skills')
EX = os.path.join(ROOT, 'examples')


def _load(name, rel):
    spec = importlib.util.spec_from_file_location(name, os.path.join(SK, rel))
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


A = _load('analyze', 'ai-visibility-check/scripts/analyze.py')
M = _load('make_outputs', 'prompt-finder/scripts/make_outputs.py')

COLS = ['run_date', 'run_label', 'project_brand', 'prompt', 'topic', 'branded', 'stage', 'platform', 'run_no',
        'status', 'brand_mentioned', 'first_position', 'brands_mentioned', 'our_domain_cited', 'cited_domains', 'answer_text']


def _row(prompt, branded='N', mentioned='N', platform='chatgpt', run_no=1, run_date='2026-10-18', brands='MG|Tesla'):
    return {'run_date': run_date, 'run_label': 'baseline', 'project_brand': 'BYD', 'prompt': prompt, 'topic': 't',
            'branded': branded, 'stage': 'Comparison', 'platform': platform, 'run_no': run_no, 'status': 'ok',
            'brand_mentioned': mentioned, 'first_position': 1 if mentioned == 'Y' else '',
            'brands_mentioned': ('BYD|' + brands) if mentioned == 'Y' else brands,
            'our_domain_cited': 'N', 'cited_domains': 'one2car.com', 'answer_text': '...'}


def _csv(path, rows, cols=COLS):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)


def _run(rows, cols=COLS, **kw):
    d = tempfile.mkdtemp()
    try:
        src = os.path.join(d, 'export.csv')
        _csv(src, rows, cols)
        for k in ('prompts',):
            if isinstance(kw.get(k), list):
                p = os.path.join(d, 'set.csv')
                _csv(p, kw[k], ['prompt', 'topic', 'branded', 'stage', 'source', 'score', 'why', 'prompt_id'])
                kw[k] = p
        out = os.path.join(d, 'report')
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            res = A.main(src, out=out, **kw)
        md = open(out + '.md', encoding='utf-8').read()
        js = json.load(open(out + '.json', encoding='utf-8'))
        return res, md, js, buf.getvalue()
    finally:
        shutil.rmtree(d)


# ---------- 5ก: Thai text, NFC, prompt IDs, wording drift ----------

def test_tone_marks_survive_and_nfc_match():
    s = 'ซื้อรถยนต์ไฟฟ้า แบตเตอรี่ ทำประกัน'
    n = A.norm(s)
    assert 'ซื้อ' in n and 'แบตเตอรี่' in n
    assert 'ำ' in n and 'ํ' not in n  # "ำ" kept as one char (NFKC would split it into ํ + า)
    assert A.norm('กุ่') == A.norm('กุ่')  # tone mark / lower vowel typed in either order
    assert A.norm('  ซื้อ​รถ   EV ') == A.norm('ซื้อรถ ev')  # zero-width + extra spaces + case
    assert unicodedata.normalize('NFKC', 'ทำ') != 'ทำ'  # why we never use NFKC


def test_same_prompt_different_encoding_gets_same_id_without_warning():
    ref = 'อยู่คอนโด ชาร์จที่บ้านไม่ได้ งบไม่เกิน 1 ล้าน รถยนต์ไฟฟ้ารุ่นไหนเหมาะที่สุด'
    variant = unicodedata.normalize('NFD', ref).replace(' ', '  ')
    rows = [_row(ref, run_no=1), _row(variant, run_no=2)]
    res, md, js, out = _run(rows, prompts=[{'prompt': ref, 'branded': 'N', 'prompt_id': 'P04'}])
    ids = {p['prompt_id'] for p in res['non_branded']['prompts']}
    assert ids == {'P04'}
    assert not [w for w in res['warnings'] if 'ถ้อยคำ' in w]


def test_wording_drift_warns_at_top_and_still_matches():
    good = 'ซื้อรถยนต์ไฟฟ้าคันแรก แบตเตอรี่ยี่ห้อไหนทนที่สุด รับประกันนานสุด'
    bad = 'ซือรถยนต์ไฟฟ้าคันแรก แบตเตอรียี่ห้อไหนทนที่สุด รับประกันนานสุด'  # tone marks lost
    other = 'รถยนต์ไฟฟ้าราคาไม่เกิน 8 แสน ยี่ห้อไหนคุ้มที่สุด'
    rows = [_row(good, run_date='2026-10-18'), _row(other, run_date='2026-10-18'),
            _row(bad, run_date='2026-10-22'), _row(other, run_date='2026-10-22')]
    res, md, js, out = _run(rows)  # no --prompts → first run is the reference
    drift = [w for w in js['warnings'] if 'ไม่ตรงชุดเดิม' in w]
    assert drift and bad in drift[0] and 'P01' in drift[0]
    assert md.index('ต้องดูก่อนอ่านผล') < md.index('## Non-branded')
    assert out.splitlines()[0].startswith('WARN:') and 'ไม่ตรงชุดเดิม' in out
    p = [e for e in res['non_branded']['prompts'] if e['prompt_id'] == 'P01'][0]
    assert p['prompt'] == good and p['platforms']['ChatGPT']['mentioned'] == '0/2'


def test_unrelated_new_prompt_gets_new_id_with_warning():
    rows = [_row('รถยนต์ไฟฟ้าราคาไม่เกิน 8 แสน ยี่ห้อไหนคุ้มที่สุด', run_date='2026-10-18'),
            _row('ชาร์จรถไฟฟ้าที่คอนโดต้องขออนุญาตนิติไหม', run_date='2026-10-22')]
    res, md, js, out = _run(rows)
    assert {e['prompt_id'] for e in res['non_branded']['prompts']} == {'P01', 'P02'}
    assert any('รหัสใหม่ P02' in w for w in res['warnings'])


# ---------- 6: branded label vs prompt text ----------

NB0 = 'รถยนต์ไฟฟ้าราคาไม่เกิน 8 แสน ยี่ห้อไหนคุ้มที่สุด'


def test_non_branded_prompt_labelled_branded_is_excluded():
    wrong = 'รถยนต์ไฟฟ้ายี่ห้อไหนรับประกันแบตนานที่สุด'
    rows = [_row(wrong, branded='Y', mentioned=m, run_no=i) for i, m in enumerate('YNN', 1)]
    rows += [_row(NB0, branded='N', mentioned='N', run_no=i) for i in (1, 2, 3)]
    res, md, js, out = _run(rows, brand='BYD', aliases='บีวายดี')
    assert res['branded']['answers'] == 0 and res['branded']['prompts'] == []
    assert [e['prompt'] for e in js['excluded_prompts']] == [wrong]
    assert res['non_branded']['answers'] == 3
    assert 'ยังไม่นับ' in md.split('## Non-branded')[0] and 'WARN:' in out


def test_branded_prompt_labelled_non_branded_does_not_inflate_rate():
    wrong = 'BYD Dolphin คุ้มกว่า MG ไหม'
    rows = [_row(wrong, branded='N', mentioned='Y', run_no=i) for i in (1, 2, 3)]
    rows += [_row(NB0, branded='N', mentioned='N', run_no=i) for i in (1, 2, 3)]
    res, md, js, out = _run(rows, brand='BYD')
    assert res['non_branded']['mention_rate'] == 0  # not 50
    assert res['non_branded']['answers'] == 3
    assert js['excluded_prompts'][0]['reason'].startswith('ติดป้าย non-branded')


def test_missing_branded_column_is_computed_from_text():
    cols = [c for c in COLS if c != 'branded']
    rows = [_row('บีวายดี Atto 3 ราคาเท่าไร', mentioned='Y'), _row(NB0, mentioned='N')]
    res, md, js, out = _run(rows, cols=cols, brand='BYD', aliases='บีวายดี, byd')
    assert res['branded']['answers'] == 1 and res['branded']['mention_rate'] == 100
    assert res['non_branded']['answers'] == 1 and res['non_branded']['mention_rate'] == 0
    assert any('คำนวณจากชื่อแบรนด์' in w for w in res['warnings'])


def test_label_values_are_normalised():
    for v in ('Y', 'yes', 'True', 'Branded', 'ใช่', ' b '):
        assert A.label(v) == 'Y', v
    for v in ('N', 'no', 'false', 'Non-branded', 'NON BRANDED', 'ไม่ใช่', 'non–branded'):
        assert A.label(v) == 'N', v
    assert A.label('') is None and A.label('maybe') is None


def test_latin_alias_matches_whole_word_only():
    names = A.brand_names('MG', [], [])
    assert A.has_brand('เทียบ MG กับ BYD', names)
    assert not A.has_brand('image quality รถไฟฟ้า', names)
    assert A.has_brand('รถBYDดีไหม', A.brand_names('BYD', [], []))


def test_no_brand_brain_is_a_note_not_a_warning():
    # Workshop 2 runs before Brand Brain (Workshop 4): note at the end, never the ⚠️ box or the next step
    rows = [_row('BYD Atto 3 ราคาเท่าไร', branded='Y', mentioned='Y')]
    res, md, js, out = _run(rows, brand='BYD')
    assert any('ยังไม่มี Brand Brain' in n and 'ต้องตรวจ' in n for n in res['notes'])
    assert not any('Brand Brain' in w for w in res['warnings'])
    assert 'Workshop 3' in md and md.index('### หมายเหตุ') > md.index('## เว็บที่ AI อ้างบ่อย')
    d = tempfile.mkdtemp()
    brain = os.path.join(d, 'brand-brain-BYD.md')
    open(brain, 'w').write('# Brand Brain')
    res, *_ = _run(rows, brand='BYD', brain=brain)
    shutil.rmtree(d)
    assert not res['notes']


# ---------- 5ข / 5ค: scoring rules in make_outputs.py ----------

def _p(prompt, source, score_given=None, demand=3, topic='t', stage='Comparison', selected=True, branded=False):
    p = {'prompt': prompt, 'topic': topic, 'stage': stage, 'branded': branded, 'source': source, 'selected': selected,
         'scores': {'business': 3, 'demand': demand, 'mention': 3, 'win': 2, 'risk': 2}, 'why': '-'}
    if score_given is not None:
        p['score'] = score_given
    return p


def _make(prompts):
    d = tempfile.mkdtemp()
    src = os.path.join(d, 'p.json')
    json.dump({'brand': 'BYD', 'aliases': ['BYD'], 'prompts': prompts}, open(src, 'w', encoding='utf-8'), ensure_ascii=False)
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        M.main(src, 'X', d)
    with open(os.path.join(d, 'prompt-set-import-X.csv'), encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    shutil.rmtree(d)
    return rows, buf.getvalue()


def test_source_ai_demand_forced_to_1():
    ai = _p('a', 'ai', demand=3)
    assert M.used_scores(ai)['demand'] == 1
    assert M.score(ai) == M.score(_p('a', 'customer', demand=1))
    assert M.score(_p('c', 'customer', demand=3)) > M.score(ai)


def test_score_mismatch_warns_and_keeps_computed():
    rows, out = _make([_p('ข้อหนึ่ง', 'customer', score_given=3.0), _p('ข้อสอง', 'customer', score_given=2.7)])
    assert 'ไม่ตรงกับที่คำนวณ' in out and 'ข้อหนึ่ง' in out and 'ข้อสอง' not in out.split('ไม่ตรงกับที่คำนวณ')[1]
    assert {r['prompt']: r['score'] for r in rows}['ข้อหนึ่ง'] == '2.7'


def test_tie_break_by_source():
    same = dict(demand=1)
    rows, _ = _make([_p('จาก AI', 'ai', **same), _p('จาก keyword', 'keyword', **same),
                     _p('จาก GSC', 'gsc', **same), _p('จากลูกค้า', 'customer', **same)])
    assert [r['prompt'] for r in rows] == ['จากลูกค้า', 'จาก GSC', 'จาก keyword', 'จาก AI']
    assert [r['prompt_id'] for r in rows] == ['P01', 'P02', 'P03', 'P04']
    assert list(rows[0].keys())[:7] == ['prompt', 'topic', 'branded', 'stage', 'source', 'score', 'why']


def test_flag_ai_selected_over_customer_same_topic_stage():
    _, out = _make([_p('AI คิด', 'ai', demand=1), _p('ลูกค้าถาม', 'customer', demand=1, selected=False),
                    _p('ลูกค้าหัวข้ออื่น', 'customer', topic='อื่น', selected=False)])
    flags = [l for l in out.splitlines() if 'AI ช่วยคิด' in l]
    assert len(flags) == 1 and 'ลูกค้าถาม' in flags[0] and 'ลูกค้าหัวข้ออื่น' not in flags[0]


# ---------- end-to-end on examples/ ----------

def test_end_to_end_examples():
    d = tempfile.mkdtemp()
    try:
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            M.main(os.path.join(EX, 'byd-prompts.json'), 'BYD', d)
        set_csv = os.path.join(d, 'prompt-set-import-BYD.csv')
        with open(set_csv, encoding='utf-8-sig') as f:
            ps = list(csv.DictReader(f))
        assert len(ps) == 10 and all(p['prompt_id'] for p in ps)
        rows = []
        for i, p in enumerate(ps):
            for plat in A.ORDER:
                for n in (1, 2, 3):
                    r = _row(p['prompt'], branded=p['branded'], mentioned='Y' if (i + n) % 3 == 0 else 'N',
                             platform=plat, run_no=n)
                    r['topic'], r['stage'] = p['topic'], p['stage']
                    rows.append(r)
        res, md, js, out = _run(rows, brand='BYD', aliases='บีวายดี', prompts=set_csv)
        assert res['excluded_prompts'] == [] and res['non_branded']['answers'] == 7 * 12
        assert res['branded']['answers'] == 3 * 12
        assert {e['prompt_id'] for b in ('non_branded', 'branded') for e in res[b]['prompts']} == {p['prompt_id'] for p in ps}
        # Excel path: fill a few Tracking rows and read them back
        from openpyxl import load_workbook
        xl = os.path.join(d, 'prompt-set-BYD.xlsx')
        wb = load_workbook(xl)
        ws = wb['Tracking (วัดด้วยมือ)']
        for r in range(2, 14):
            ws.cell(r, 1, '2026-10-18'); ws.cell(r, 6, 'Y' if r % 2 else 'N'); ws.cell(r, 8, 'BYD, MG')
        wb.save(xl)
        with contextlib.redirect_stdout(io.StringIO()):
            res = A.main(xl, brand='BYD', out=os.path.join(d, 'x'))
        assert res['rows'] == 12 and res['excluded_prompts'] == []
        assert res['non_branded']['prompts'][0]['prompt_id'] == 'P01'
    finally:
        shutil.rmtree(d)


# ---------- class metrics (M1/M3): SoV = our names ÷ all brand names, n and ± ----------

def test_sov_class_definition_and_margin():
    # slide M1 example: 10 answers, us in 4, 16 brand names in total → Mention 40% ±30 · SoV 4/16 = 25%
    brands = ['BYD|MG', 'MG', 'BYD|MG', 'MG|Tesla', 'Tesla', 'BYD|MG', 'MG', 'Tesla', 'BYD|Tesla|MG', 'MG']
    assert sum(len(b.split('|')) for b in brands) == 16
    rows = []
    for i, b in enumerate(brands, 1):
        r = _row(NB0, mentioned='Y' if b.startswith('BYD') else 'N', run_no=i)
        r['brands_mentioned'] = b + ('|BYD' if b.startswith('BYD') else '')  # duplicate name in one answer counts once
        rows.append(r)
    res, md, js, out = _run(rows, brand='BYD')
    nb = res['non_branded']
    assert nb['mention_rate'] == 40 and nb['answers'] == 10 and nb['margin'] == 30
    assert nb['sov_all'] == 25.0 and nb['brands_named'] == 3
    assert nb['sov_band'].startswith('ต่ำกว่าค่าเฉลี่ย 33.3%')  # 3 brands → bands for 8–15 brands don't apply
    assert A.margin(40, 120) == 9
    assert [A.sov_band(v, 12) for v in (4.9, 5, 10, 20)] == ['แทบไม่มีเสียง', 'อยู่ในวง', 'กลุ่มนำ', 'เจ้าที่']
    assert '±30 (n = 10' in md and 'ถ่วงตำแหน่ง' not in md
